"""Self-check of the export formats, no database needed.

Builds one report with every column kind (including Latvian letters and an
empty table) and checks that each format opens and carries the same rows:

    docker compose -p marketpulse exec -T app python scripts/check_exports.py
"""
from __future__ import annotations

import csv
import io
import json
import sys
import zipfile
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl  # noqa: E402

from app.api import exports as E  # noqa: E402

ROWS = [
    ["Hermès", "Terre d'Hermès", 89.5, -12.3, 4.2, 3, True, datetime(2026, 10, 1, 8, 30), "https://example.lv/a"],
    ["Ajmal", "Nacionālais dvēseles uzmundrinājums", 20.03, 69.0, 0.0, 1, False, datetime(2026, 10, 2), None],
]
COLS = [("Brand", "text"), ("Perfume", "text"), ("Price", "eur"), ("Change", "pct"), ("Gap", "share"),
        ("Shops", "int"), ("In stock", "bool"), ("Checked", "datetime"), ("URL", "url")]


def report() -> E._Report:
    index = E._Table("Price index", [("Date", "date"), ("Notino", "num"), ("Douglas", "num")],
                     [["2026-10-01", 100.0, 100.0], ["2026-10-02", 101.5, 99.2]], chart=("line", [1, 2]))
    return E._Report("Check", "Filters: none", [
        E._Table("Listings", COLS, ROWS, note="two rows"),
        E._Table("Empty", COLS, []),
        index,
    ], kpis=[("Listings", "2"), ("Median gap", "4.2%"), ("Change", "−12.3%", "down")], notes=["A note."])


def check() -> None:
    rep = report()

    wb = openpyxl.load_workbook(io.BytesIO(E._report_xlsx(rep).getvalue()))
    assert wb.sheetnames == ["Report", "Listings", "Empty", "Price index"], wb.sheetnames
    ws = wb["Listings"]
    assert list(ws.tables) and ws["A5"].value == "Hermès" and ws["C5"].value == 89.5
    assert ws["C5"].number_format == E._XL_FMT["eur"] and ws["I5"].hyperlink.target == ROWS[0][8]
    assert ws.freeze_panes == "A5" and wb["Price index"]._charts
    assert wb["Report"]["A5"].value == 2   # KPI written as a number, not as text

    pdf = E._report_pdf(rep).getvalue()
    assert pdf.startswith(b"%PDF") and b"DejaVuSans" in pdf

    buf, name, media = E._report_csv(rep, "check")
    assert name == "check.zip" and media == "application/zip"
    with zipfile.ZipFile(buf) as z:
        assert z.namelist() == ["01_listings.csv", "02_empty.csv", "03_price_index.csv", "README.txt"]
        rows = list(csv.reader(io.StringIO(z.read("01_listings.csv").decode("utf-8-sig"))))
    assert rows[0][2] == "Price (EUR)" and rows[2][1] == ROWS[1][1] and rows[1][6] == "yes"

    one = E._Report("One", "", [rep.tables[0]])
    buf, name, _ = E._report_csv(one, "one")
    assert name == "one.csv" and buf.getvalue().startswith("﻿".encode())

    doc = json.loads(E._report_json(rep).getvalue())
    assert doc["schema_version"] == E.SCHEMA_VERSION and doc["currency"] == "EUR"
    t = doc["tables"][0]
    assert t["row_count"] == 2 and t["columns"][2] == {"key": "price", "label": "Price", "type": "eur", "unit": "EUR"}
    assert t["rows"][0]["checked"] == "2026-10-01T08:30:00Z" and t["rows"][1]["in_stock"] is False
    assert doc["summary"][2]["value"] == -12.3 and doc["tables"][2]["rows"][0]["date"] == "2026-10-01"
    print("exports OK: xlsx, pdf, csv, zip, json")


if __name__ == "__main__":
    check()
