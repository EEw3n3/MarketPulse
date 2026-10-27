"""Export endpoints — Excel, PDF, CSV and JSON, generated on the server.

Exports are a Pro feature (see app/api/main.py). The catalogue export takes the
same filters as the catalogue plus an explicit selection (`keys` for perfume
cards, `ids` for shop listings); the analytics export writes every section of
the cabinet. Every export is built as one `_Report` (title, filters, key
figures, notes and a list of tables), so the four formats always carry the
same columns and the same figures:

- xlsx: a "Report" cover sheet (key figures, contents, notes), then one sheet
  per table as a native Excel table (filters, banding, frozen header, number
  formats), native charts where a table has one, print setup and page footer.
- pdf: A4 with a running header and "Page X of Y" footer, key-figure tiles,
  charts, tables that repeat their header on every page. DejaVu Sans (shipped
  with matplotlib) is embedded so Latvian and Cyrillic letters print.
- csv: one table is one UTF-8 CSV; several tables come as a ZIP with one CSV
  per table and a README. Units are in the header ("Price (EUR)").
- json: a versioned document (`schema_version`) with typed column metadata.

A single product has its own Excel summary sheet and PDF first page (photo,
price, period statistics, chart), followed by the same generic tables.
"""
from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any, List, Optional
from xml.sax.saxutils import escape

# matplotlib MUST use Agg backend before pyplot import — Docker has no display.
import matplotlib
matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter

import httpx
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.worksheet.table import Table as XLTable, TableStyleInfo
from reportlab.lib import colors as rl_colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.platypus import (
    CondPageBreak,
    Image as RLImage,
    KeepTogether,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select
from sqlalchemy.orm import selectinload

from app.api.cards import cards_for_keys, query_cards, query_listings
from app.api.deps import get_db
from app.domain.models import Product
from app.services import market
from app.services.price_series import STORE_LABELS, store_key, store_label

router = APIRouter(prefix="/exports", tags=["exports"])

_FORMATS = "^(xlsx|pdf|csv|json)$"
SCHEMA_VERSION = "2.0"
SOURCE = "marketpulse1.me"


# ── Constants ─────────────────────────────────────────────────────────────────

PERIOD_DAYS = {"1w": 7, "1m": 30, "3m": 90, "6m": 180, "1y": 365, "all": None}
PERIOD_LABEL = {"1w": "1 week", "1m": "1 month", "3m": "3 months", "6m": "6 months", "1y": "1 year", "all": "All time"}

CHANGE_LABEL = {
    "synthetic": "Estimate",
    "initial": "First seen",
    "price_down": "Price drop",
    "price_up": "Price rise",
    "no_change": "No change",
    "back_in_stock": "Back in stock",
    "out_of_stock": "Out of stock",
}

# Ink, accent and the down/up colours of the site (see frontend/src/index.css).
INK = "111315"
ACCENT = "D0106F"
GREEN = "1D7A3E"
RED = "B3262F"
GRAY = "6B7280"
MUTED = "9CA3AF"
LINE = "E5E7EB"
LIGHT = "F4F5F7"
LINK = "2563EB"

# Chart lines of each shop (frontend/src/lib/shops.ts).
SHOP_LINE = {"notino": "#4766E8", "makeup": "#D9622B", "220": "#3B8F22", "douglas": "#B98A00", "1a": "#7C3AED"}
_LABEL_TO_SHOP = {v: k for k, v in STORE_LABELS.items()}

_NUMERIC = ("eur", "pct", "share", "int", "num")


# ── Report model: one source, four formats ───────────────────────────────────


@dataclass
class _Table:
    title: str
    columns: list[tuple[str, str]]   # (label, kind): text eur pct int num bool url date datetime
    rows: list[list[Any]]
    note: str = ""                   # one line under the table title
    chart: Optional[tuple[str, list[int]]] = None   # ("line" | "bar", value column indexes); x = column 0
    heat: bool = False               # colour scale over the numeric cells

    def __post_init__(self):
        # A price of 0 means "unknown" (see CLAUDE.md), never "free": show it as empty.
        eur = [i for i, (_, k) in enumerate(self.columns) if k == "eur"]
        for r in self.rows:
            for i in eur:
                if r[i] is not None and not float(r[i]):
                    r[i] = None


@dataclass
class _Report:
    title: str                       # "Catalogue", "Market analytics", product name
    subtitle: str                    # filters / period in one line
    tables: list[_Table]
    kpis: list[tuple] = field(default_factory=list)    # (label, value text[, tone "down" | "up"])
    notes: list[str] = field(default_factory=list)
    generated: datetime = field(default_factory=datetime.utcnow)

    @property
    def stamp(self) -> str:
        return self.generated.strftime("%Y-%m-%d %H:%M UTC")


# ── Helpers ───────────────────────────────────────────────────────────────────


def _period_cutoff(period: str) -> Optional[datetime]:
    days = PERIOD_DAYS.get(period)
    if days is None:
        return None
    return datetime.utcnow() - timedelta(days=days)


def _safe_filename(s: str, max_len: int = 60) -> str:
    keep = "-_.() "
    out = "".join(c if c.isalnum() or c in keep else "_" for c in s).strip()
    return (out[:max_len] or "export").rstrip(".")


def _stream(buf: io.BytesIO, filename: str, media_type: str) -> StreamingResponse:
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _fetch_image_bytes(url: Optional[str], timeout: float = 3.0) -> Optional[bytes]:
    """Best-effort image fetch. Returns None on any failure — caller renders placeholder."""
    if not url or not isinstance(url, str) or not url.startswith(("http://", "https://")):
        return None
    try:
        with httpx.Client(timeout=timeout, follow_redirects=True) as cli:
            r = cli.get(url, headers={"User-Agent": "Mozilla/5.0 (compatible; MarketPulse/1.0)"})
            ct = r.headers.get("content-type", "")
            if r.status_code == 200 and ct.startswith("image/") and r.content:
                return _shrink_image(r.content)
    except Exception:
        pass
    return None


def _shrink_image(data: bytes, size: int = 480) -> bytes:
    """Shop photos are often 2000 px; a 480 px JPEG keeps the files small and sharp enough."""
    try:
        from PIL import Image

        im = Image.open(io.BytesIO(data))
        im.thumbnail((size, size))
        if im.mode in ("RGBA", "LA", "P"):
            im = im.convert("RGBA")
            bg = Image.new("RGB", im.size, "white")
            bg.paste(im, mask=im.split()[-1])
            im = bg
        out = io.BytesIO()
        im.convert("RGB").save(out, format="JPEG", quality=88, optimize=True)
        return out.getvalue()
    except Exception:
        return data


def _num(v) -> Optional[float]:
    return round(float(v), 2) if v is not None else None


def _delta_percent(current: Optional[float], previous: Optional[float]) -> Optional[float]:
    if not current or not previous or previous <= 0 or previous == current:
        return None
    return round((current - previous) / previous * 100, 1)


def _eur(v: Optional[float]) -> str:
    return "—" if v is None else f"{v:,.2f} €".replace(",", " ")


def _pct(v: Optional[float]) -> str:
    if v is None:
        return "—"
    return f"{v:+.1f}%".replace("-", "−") if v else "0%"


def _tone(v: Optional[float]) -> Optional[str]:
    return None if not v else ("down" if v < 0 else "up")


def _as_date(v):
    """Analytics dates arrive as ISO strings; Excel and JSON want real dates."""
    if isinstance(v, str) and len(v) == 10:
        try:
            return date.fromisoformat(v)
        except ValueError:
            return v
    return v


def _plain(v, kind: str):
    """Value as text for CSV and JSON."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if kind == "date" else v.strftime("%Y-%m-%dT%H:%M:%SZ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, Decimal):
        return float(v)
    return v


def _key(label: str) -> str:
    out = "".join(c if c.isalnum() else "_" for c in label.lower())
    while "__" in out:
        out = out.replace("__", "_")
    return out.strip("_") or "value"


def _unit(kind: str) -> Optional[str]:
    return {"eur": "EUR", "pct": "%", "share": "%"}.get(kind)


# ── Charts (PNG for PDF) ─────────────────────────────────────────────────────


def _mpl_axes(width_cm: float, height_cm: float):
    fig, ax = plt.subplots(figsize=(width_cm / 2.54, height_cm / 2.54), dpi=200)
    ax.tick_params(axis="both", labelsize=7, colors="#4B5563", length=0, pad=4)
    for spine in ("top", "right", "left"):
        ax.spines[spine].set_visible(False)
    ax.spines["bottom"].set_color("#D1D5DB")
    ax.grid(True, axis="y", color="#E5E7EB", linewidth=0.6)
    ax.set_axisbelow(True)
    return fig, ax


def _png(fig) -> io.BytesIO:
    fig.tight_layout(pad=0.4)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


def _history_png(points: list[tuple[datetime, float]], width_cm: float, height_cm: float) -> io.BytesIO:
    """Step line of a product's price: the price holds until the next check that changed it."""
    fig, ax = _mpl_axes(width_cm, height_cm)
    if len(points) >= 1:
        xs = [p[0] for p in points]
        ys = [p[1] for p in points]
        if len(points) == 1 or xs[-1] < datetime.utcnow() - timedelta(hours=1):
            xs.append(datetime.utcnow())
            ys.append(ys[-1])
        ax.step(xs, ys, where="post", color=f"#{INK}", linewidth=1.6)
        ax.fill_between(xs, ys, min(ys) * 0.97, step="post", color=f"#{ACCENT}", alpha=0.07)
        lo, hi = min(ys), max(ys)
        pad = max((hi - lo) * 0.25, hi * 0.03)
        ax.set_ylim(lo - pad, hi + pad)
        i_lo, i_hi = ys.index(lo), ys.index(hi)
        if hi > lo:
            ax.scatter([xs[i_lo]], [lo], color=f"#{GREEN}", s=14, zorder=3)
            ax.annotate(f"min {_eur(lo)}", (xs[i_lo], lo), textcoords="offset points", xytext=(4, -10),
                        fontsize=6.5, color=f"#{GREEN}")
            ax.scatter([xs[i_hi]], [hi], color=f"#{RED}", s=14, zorder=3)
            ax.annotate(f"max {_eur(hi)}", (xs[i_hi], hi), textcoords="offset points", xytext=(4, 5),
                        fontsize=6.5, color=f"#{RED}")
        ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:,.0f} €"))
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%d %b"))
        ax.xaxis.set_major_locator(mdates.AutoDateLocator(maxticks=8))
    else:
        ax.text(0.5, 0.5, "No price records in this period", ha="center", va="center",
                transform=ax.transAxes, color=f"#{GRAY}", fontsize=8)
        ax.set_xticks([])
        ax.set_yticks([])
    return _png(fig)


def _table_chart_png(t: _Table, width_cm: float, height_cm: float) -> Optional[io.BytesIO]:
    if not t.chart or not t.rows:
        return None
    kind, cols = t.chart
    xs = [_as_date(r[0]) for r in t.rows]
    fig, ax = _mpl_axes(width_cm, height_cm)
    if kind == "line":
        for ci in cols:
            label = t.columns[ci][0]
            ys = [r[ci] for r in t.rows]
            ax.plot(xs, ys, label=label, linewidth=1.5,
                    color=SHOP_LINE.get(_LABEL_TO_SHOP.get(label, ""), f"#{INK}"))
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%d %b"))
        ax.legend(fontsize=7, frameon=False, ncol=len(cols), loc="upper left")
    else:
        n = len(cols)
        w = 0.8 / n
        palette = [f"#{INK}", f"#{ACCENT}", "#9CA3AF"]
        for j, ci in enumerate(cols):
            ys = [r[ci] or 0 for r in t.rows]
            pos = [i + (j - (n - 1) / 2) * w for i in range(len(xs))]
            ax.bar(pos, ys, width=w * 0.92, label=t.columns[ci][0], color=palette[j % len(palette)])
        ax.set_xticks(range(len(xs)))
        ax.set_xticklabels([str(x) for x in xs])
        ax.legend(fontsize=7, frameon=False, ncol=n, loc="upper right")
    return _png(fig)


def _placeholder_png(brand: Optional[str]) -> io.BytesIO:
    """Light square with the brand's initial — when the shop's photo can't be fetched."""
    initial = (brand or "?").strip()[:1].upper() or "?"
    fig, ax = plt.subplots(figsize=(2, 2), dpi=100)
    fig.patch.set_facecolor(f"#{LIGHT}")
    ax.set_facecolor(f"#{LIGHT}")
    ax.text(0.5, 0.5, initial, ha="center", va="center", color=f"#{MUTED}", fontsize=72, fontweight="bold")
    ax.axis("off")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor=f"#{LIGHT}", pad_inches=0)
    plt.close(fig)
    buf.seek(0)
    return buf


# ── Excel ────────────────────────────────────────────────────────────────────

_XL_FMT = {
    "eur": '#,##0.00 "€"',
    "pct": '+0.0"%";"−"0.0"%";0"%"',
    "share": '0.0"%"',
    "int": "#,##0",
    "num": "0.00",
    "date": "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm",
}
_HAIR = Side(border_style="thin", color=LINE)


def _xl_text_width(v, kind: str) -> int:
    if v is None:
        return 1
    if kind == "eur":
        return len(f"{v:,.2f} €")
    if kind in ("pct", "share"):
        return 7
    if kind == "datetime":
        return 16
    if kind == "date":
        return 10
    if kind == "url":
        return 12
    return len(str(v))


def _xl_number(text):
    """"25,878" -> (25878, "#,##0"), "22.9%" -> (22.9, '0.0"%"'); anything else -> (text, None).
    Excel flags numbers stored as text with a green corner."""
    if not isinstance(text, str):
        return text, None
    raw = text.replace(",", "").replace(" ", "").strip()
    try:
        if raw.endswith("%"):
            return float(raw[:-1]), '0.0"%"'
        return (int(raw), "#,##0") if raw.lstrip("-").isdigit() else (text, None)
    except ValueError:
        return text, None


def _xl_head(ws, title: str, line: str) -> None:
    ws["A1"] = title
    ws["A1"].font = Font(name="Calibri", size=15, bold=True, color=INK)
    ws.row_dimensions[1].height = 24
    ws["A2"] = line
    ws["A2"].font = Font(name="Calibri", size=9, color=GRAY)


def _xl_print(ws, header_row: Optional[int], title: str, landscape_page: bool = True) -> None:
    ws.page_setup.orientation = "landscape" if landscape_page else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    if header_row:
        ws.print_title_rows = f"{header_row}:{header_row}"
    ws.oddHeader.left.text = f"MarketPulse · {title}"
    ws.oddHeader.left.size = 8
    ws.oddFooter.left.text = SOURCE
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8


def _xl_table(ws, t: _Table, top: int, name: str) -> int:
    """Write `t` as a native Excel table with its header on row `top`. Returns the last row."""
    ncol = len(t.columns)
    for ci, (label, kind) in enumerate(t.columns, start=1):
        c = ws.cell(row=top, column=ci, value=label)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=INK)
        # indent keeps right-aligned labels clear of the filter arrow
        c.alignment = Alignment(horizontal="right", indent=2, vertical="center") if kind in _NUMERIC             else Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[top].height = 20
    label_of = {ci: label for ci, (label, _) in enumerate(t.columns, start=1)}

    for ri, row in enumerate(t.rows, start=top + 1):
        for ci, (v, (_, kind)) in enumerate(zip(row, t.columns), start=1):
            if kind == "bool" and v is not None:
                v = "Yes" if v else "No"
            if kind == "date":
                v = _as_date(v)
            if kind == "url" and v:
                c = ws.cell(row=ri, column=ci, value="Open ↗")
                c.hyperlink = v
                c.font = Font(color=LINK, underline="single")
                if (ri - top) % 2 == 0:
                    c.fill = PatternFill("solid", fgColor=LIGHT)
                continue
            fmt = _XL_FMT.get(kind)
            if label_of[ci] == "Value":
                v, fmt = _xl_number(v)
            c = ws.cell(row=ri, column=ci, value=v)
            if fmt:
                c.number_format = fmt
            if kind == "bool" or (label_of[ci] == "Value" and fmt):
                c.alignment = Alignment(horizontal="center" if kind == "bool" else "left")
            if (ri - top) % 2 == 0:
                c.fill = PatternFill("solid", fgColor=LIGHT)

    last = top + len(t.rows)
    end_col = get_column_letter(ncol)
    if t.rows:
        xt = XLTable(displayName=name, ref=f"A{top}:{end_col}{last}")
        xt.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(xt)
        for ci, (_, kind) in enumerate(t.columns, start=1):
            col = get_column_letter(ci)
            rng = f"{col}{top + 1}:{col}{last}"
            if kind == "pct" and not t.heat:
                ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"],
                                                              font=Font(color=GREEN, bold=True)))
                ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"],
                                                              font=Font(color=RED, bold=True)))
        if t.heat:
            first_num = next((i for i, (_, k) in enumerate(t.columns, start=1) if k in _NUMERIC), None)
            if first_num:
                ws.conditional_formatting.add(
                    f"{get_column_letter(first_num)}{top + 1}:{end_col}{last}",
                    ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                   end_type="max", end_color="F3A6C8"))
    else:
        last = top + 1
        ws.cell(row=last, column=1, value="No rows for these filters.").font = Font(italic=True, color=GRAY)

    for ci, (label, kind) in enumerate(t.columns, start=1):
        content = max([_xl_text_width(r[ci - 1], kind) for r in t.rows[:500]] + [0]) + 3
        header = len(label) * 1.2 + 6   # bold text plus the filter arrow
        limit = 46 if kind == "text" else 24
        ws.column_dimensions[get_column_letter(ci)].width = max(9, min(limit, max(content, header)))
    ws.freeze_panes = f"A{top + 1}"
    ws.sheet_view.showGridLines = False
    return last


def _xl_chart(ws, t: _Table, top: int, last: int) -> None:
    if not t.chart or len(t.rows) < 2:
        return
    kind, cols = t.chart
    ch = LineChart() if kind == "line" else BarChart()
    ch.height, ch.width = 9, 22
    ch.y_axis.delete = False
    ch.x_axis.delete = False
    for ci in cols:
        ch.add_data(Reference(ws, min_col=ci + 1, min_row=top, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=last))
    if kind == "line":
        ch.x_axis.number_format = "dd.mm"
        _xl_fit_axis(ch, [r[ci] for r in t.rows for ci in cols])
        for s, ci in zip(ch.series, cols):
            color = SHOP_LINE.get(_LABEL_TO_SHOP.get(t.columns[ci][0], ""), f"#{INK}").lstrip("#")
            s.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line.width = 22000
            s.smooth = False
    else:
        for s, color in zip(ch.series, [INK, ACCENT, MUTED]):
            s.graphicalProperties.solidFill = color
            s.graphicalProperties.line.solidFill = color
    ch.legend.position = "t"
    ws.add_chart(ch, f"{get_column_letter(len(t.columns) + 2)}{top}")


def _xl_fit_axis(ch, values: list) -> None:
    """Start the value axis near the data instead of at zero, so small moves are visible."""
    nums = [v for v in values if isinstance(v, (int, float))]
    if not nums:
        return
    lo, hi = min(nums), max(nums)
    pad = max((hi - lo) * 0.15, abs(hi) * 0.01, 0.5)
    ch.y_axis.scaling.min = round(lo - pad)
    ch.y_axis.scaling.max = round(hi + pad)


def _sheet_names(tables: list[_Table]) -> list[str]:
    names, used = [], {"Report"}
    for t in tables:
        name = "".join(c for c in t.title if c not in "[]:*?/\\")[:31].strip() or "Sheet"
        base, n = name, 2
        while name in used:
            name = f"{base[:28]} {n}"
            n += 1
        used.add(name)
        names.append(name)
    return names


def _xl_cover(ws, rep: _Report, names: list[str], start_row: int = 1) -> int:
    """Key figures, contents with links to every sheet, and notes. Returns the next free row."""
    r = start_row
    if rep.kpis:
        for i, k in enumerate(rep.kpis[:6]):
            col = 1 + i
            lab = ws.cell(row=r, column=col, value=k[0].upper())
            lab.font = Font(size=8, bold=True, color=GRAY)
            num, fmt = _xl_number(k[1])
            val = ws.cell(row=r + 1, column=col, value=num)
            if fmt:
                val.number_format = fmt
            tone = k[2] if len(k) > 2 else None
            val.font = Font(size=16, bold=True, color=GREEN if tone == "down" else RED if tone == "up" else INK)
            for rr in (r, r + 1):
                cell = ws.cell(row=rr, column=col)
                cell.fill = PatternFill("solid", fgColor=LIGHT)
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 18
        ws.row_dimensions[r + 1].height = 28
        r += 3

    if names:
        ws.cell(row=r, column=1, value="Contents").font = Font(size=12, bold=True, color=INK)
        r += 1
        for name, t in zip(names, rep.tables):
            c = ws.cell(row=r, column=1, value=t.title)
            c.hyperlink = f"#{quote_sheetname(name)}!A1"
            c.font = Font(color=LINK, underline="single")
            n = ws.cell(row=r, column=3, value=f"{len(t.rows):,} rows")
            n.font = Font(color=GRAY)
            for ci in range(1, 7):
                ws.cell(row=r, column=ci).border = Border(bottom=_HAIR)
            r += 1
        r += 1

    if rep.notes:
        ws.cell(row=r, column=1, value="About this data").font = Font(size=12, bold=True, color=INK)
        r += 1
        for note in rep.notes:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            c = ws.cell(row=r, column=1, value=f"•  {note}")
            c.font = Font(size=10, color="374151")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 15 * (1 + len(note) // 105)
            r += 1
    return r


def _xl_workbook(rep: _Report):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.properties.title = f"MarketPulse · {rep.title}"
    wb.properties.subject = rep.subtitle
    wb.properties.creator = "MarketPulse"
    wb.properties.keywords = "perfume prices, Latvia"
    return wb


def _report_xlsx(rep: _Report) -> io.BytesIO:
    wb = _xl_workbook(rep)
    names = _sheet_names(rep.tables)

    cover = wb.active
    cover.title = "Report"
    _xl_head(cover, f"MarketPulse · {rep.title}", f"{rep.subtitle}  ·  generated {rep.stamp}")
    _xl_cover(cover, rep, names, start_row=4)
    for col in "ABCDEF":
        cover.column_dimensions[col].width = 21
    cover.sheet_view.showGridLines = False
    _xl_print(cover, None, rep.title, landscape_page=False)

    for i, (t, name) in enumerate(zip(rep.tables, names), start=1):
        ws = wb.create_sheet(name)
        _xl_head(ws, t.title, "  ·  ".join(x for x in (t.note, rep.subtitle, f"generated {rep.stamp}") if x))
        last = _xl_table(ws, t, 4, f"tbl{i}_{_key(t.title)[:40]}")
        _xl_chart(ws, t, 4, last)
        _xl_print(ws, 4, rep.title)

    buf = io.BytesIO()
    wb.save(buf)
    return buf


# ── PDF ──────────────────────────────────────────────────────────────────────

_FONT_DIR = Path(matplotlib.get_data_path()) / "fonts" / "ttf"


def _register_fonts() -> tuple[str, str]:
    """DejaVu Sans covers Latvian (ā, ē, š, ž) and Cyrillic; Helvetica does not."""
    try:
        if "MP" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("MP", str(_FONT_DIR / "DejaVuSans.ttf")))
            pdfmetrics.registerFont(TTFont("MP-Bold", str(_FONT_DIR / "DejaVuSans-Bold.ttf")))
            pdfmetrics.registerFont(TTFont("MP-Italic", str(_FONT_DIR / "DejaVuSans-Oblique.ttf")))
            pdfmetrics.registerFont(TTFont("MP-BoldItalic", str(_FONT_DIR / "DejaVuSans-BoldOblique.ttf")))
            pdfmetrics.registerFontFamily("MP", normal="MP", bold="MP-Bold", italic="MP-Italic",
                                          boldItalic="MP-BoldItalic")
        return "MP", "MP-Bold"
    except Exception:
        return "Helvetica", "Helvetica-Bold"


FONT, FONT_BOLD = _register_fonts()


def _hex(c: str):
    return rl_colors.HexColor(f"#{c.lstrip('#')}")


class _St:
    """Paragraph styles of every PDF."""
    title = ParagraphStyle("t", fontName=FONT_BOLD, fontSize=19, leading=23, textColor=_hex(INK))
    sub = ParagraphStyle("s", fontName=FONT, fontSize=8.5, leading=12, textColor=_hex(GRAY), spaceAfter=10)
    h2 = ParagraphStyle("h2", fontName=FONT_BOLD, fontSize=11.5, leading=14, textColor=_hex(INK),
                        spaceBefore=12, spaceAfter=2)
    note = ParagraphStyle("n", fontName=FONT, fontSize=7.5, leading=10, textColor=_hex(GRAY), spaceAfter=5)
    body = ParagraphStyle("b", fontName=FONT, fontSize=9, leading=13, textColor=_hex("1F2937"))
    bullet = ParagraphStyle("bl", parent=body, fontSize=8, leading=11.5, leftIndent=10, bulletIndent=0,
                            textColor=_hex("374151"))
    cell = ParagraphStyle("c", fontName=FONT, fontSize=7.5, leading=9.2, textColor=_hex("1F2937"))
    kpi_label = ParagraphStyle("kl", fontName=FONT_BOLD, fontSize=6.5, leading=8, textColor=_hex(GRAY))
    kpi_value = ParagraphStyle("kv", fontName=FONT_BOLD, fontSize=14, leading=17, textColor=_hex(INK))


def _canvas_maker(title: str, stamp: str):
    """Canvas that draws the running header and a "Page X of Y" footer once the page count is known."""

    class _Canvas(pdfcanvas.Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._pages: list[dict] = []

        def showPage(self):
            self._pages.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._pages)
            for state in self._pages:
                self.__dict__.update(state)
                self._chrome(total)
                pdfcanvas.Canvas.showPage(self)
            pdfcanvas.Canvas.save(self)

        def _chrome(self, total: int) -> None:
            w, h = self._pagesize
            m = 1.5 * cm
            self.saveState()
            self.setFillColor(_hex(ACCENT))
            self.circle(m + 2.2, h - 1.05 * cm + 2.6, 2.6, stroke=0, fill=1)
            self.setFillColor(_hex(INK))
            self.setFont(FONT_BOLD, 9.5)
            self.drawString(m + 9, h - 1.05 * cm, "MarketPulse")
            self.setFont(FONT, 8)
            self.setFillColor(_hex(GRAY))
            self.drawRightString(w - m, h - 1.05 * cm, title[:90])
            self.setStrokeColor(_hex(LINE))
            self.setLineWidth(0.6)
            self.line(m, h - 1.3 * cm, w - m, h - 1.3 * cm)
            self.line(m, 1.25 * cm, w - m, 1.25 * cm)
            self.setFont(FONT, 7)
            self.drawString(m, 0.85 * cm, f"{SOURCE}  ·  prices in EUR as listed by the shops  ·  generated {stamp}")
            self.drawRightString(w - m, 0.85 * cm, f"Page {self._pageNumber} of {total}")
            self.restoreState()

    return _Canvas


def _pdf_doc(buf, page, title: str) -> SimpleDocTemplate:
    return SimpleDocTemplate(
        buf, pagesize=page, topMargin=1.9 * cm, bottomMargin=1.7 * cm,
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        title=f"MarketPulse · {title}", author="MarketPulse", subject=title, creator=SOURCE,
    )


def _pdf_kpis(kpis: list[tuple], width: float) -> Table:
    n = len(kpis)
    gap = 0.25 * cm
    w = (width - gap * (n - 1)) / n
    cells, widths, cmds = [], [], [("VALIGN", (0, 0), (-1, -1), "TOP")]
    for i, k in enumerate(kpis):
        if i:
            cells.append("")
            widths.append(gap)
        tone = k[2] if len(k) > 2 else None
        color = GREEN if tone == "down" else RED if tone == "up" else INK
        val = ParagraphStyle("kv2", parent=_St.kpi_value, textColor=_hex(color),
                             fontSize=14 if len(k[1]) < 12 else 11)
        cells.append([Paragraph(escape(k[0].upper()), _St.kpi_label), Spacer(1, 3), Paragraph(escape(k[1]), val)])
        widths.append(w)
        col = len(cells) - 1
        cmds += [("BACKGROUND", (col, 0), (col, 0), _hex(LIGHT)),
                 ("LEFTPADDING", (col, 0), (col, 0), 9), ("TOPPADDING", (col, 0), (col, 0), 8),
                 ("BOTTOMPADDING", (col, 0), (col, 0), 9)]
    t = Table([cells], colWidths=widths)
    t.setStyle(TableStyle(cmds))
    return t


def _pdf_cell_text(v, kind: str) -> str:
    if v is None:
        return "—"
    if kind == "eur":
        return _eur(v)
    if kind == "pct":
        return _pct(v)
    if kind == "share":
        return f"{v:.1f}%"
    if kind == "int":
        return f"{v:,}".replace(",", " ")
    if kind == "num":
        return f"{v:.1f}"
    if kind == "bool":
        return "Yes" if v else "No"
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if kind == "date" else v.strftime("%Y-%m-%d %H:%M")
    return str(v)


def _pdf_table(t: _Table, usable: float) -> Table:
    url_i = next((i for i, (_, k) in enumerate(t.columns) if k == "url"), None)
    link_i = next((i for i, (c, _) in enumerate(t.columns) if c in ("Perfume", "Title")), None)
    shown = [(i, c, k) for i, (c, k) in enumerate(t.columns) if k != "url"]

    # Width from content. Numbers and short codes keep their natural width; text
    # columns share the rest but never get narrower than their longest word.
    def span(s: str, bold: bool = False) -> float:
        return pdfmetrics.stringWidth(s, FONT_BOLD if bold else FONT, 7.5) + 12

    nat, low = [], []
    for i, label, kind in shown:
        texts = [_pdf_cell_text(r[i], kind) for r in t.rows[:300]]
        widest = max([span(x[:42]) for x in texts] + [0])
        word = max([span(w) for x in texts for w in x.split()] + [span(w, True) for w in label.split()] + [12])
        if kind == "text":
            nat.append(max(widest, span(label, True)))
            low.append(word)
        else:
            nat.append(max(widest, word))
            low.append(nat[-1])
    text_i = [j for j, (_, _, k) in enumerate(shown) if k == "text"]
    fixed = sum(w for j, w in enumerate(nat) if j not in text_i)
    room = usable - fixed
    text_nat = sum(nat[j] for j in text_i)
    widths = list(nat)
    narrow = sum(nat) < usable * 0.55   # a small table keeps its own width instead of stretching
    if narrow:
        widths = [w + (18 if j in text_i else 6) for j, w in enumerate(nat)]
    elif text_i and text_nat <= room:
        for j in text_i:
            widths[j] = nat[j] * room / text_nat
    elif text_i:
        text_low = sum(low[j] for j in text_i)
        f = max(0.0, (room - text_low) / (text_nat - text_low)) if text_nat > text_low else 0.0
        for j in text_i:
            widths[j] = low[j] + (nat[j] - low[j]) * f
    if sum(widths) > usable:
        widths = [w * usable / sum(widths) for w in widths]

    head = [Paragraph(f'<font color="white"><b>{escape(c)}</b></font>',
                      ParagraphStyle("h", parent=_St.cell, alignment=2 if k in _NUMERIC else 0))
            for _, c, k in shown]
    data = [head]
    for r in t.rows:
        line = []
        for i, _, kind in shown:
            txt = _pdf_cell_text(r[i], kind)
            if kind == "text":
                txt = escape(txt)
                if i == link_i and url_i is not None and r[url_i]:
                    txt = f'<link href="{escape(r[url_i]).replace(chr(34), "%22")}">{txt}</link>'
                line.append(Paragraph(txt, _St.cell))
            else:
                line.append(txt)
        data.append(line)
    if len(data) == 1:
        data.append([Paragraph("<i>No rows for these filters.</i>", _St.cell)] + [""] * (len(shown) - 1))

    cmds = [
        ("BACKGROUND", (0, 0), (-1, 0), _hex(INK)),
        ("FONTNAME", (0, 1), (-1, -1), FONT),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("TEXTCOLOR", (0, 1), (-1, -1), _hex("1F2937")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [rl_colors.white, _hex(LIGHT)]),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, _hex(LINE)),
        ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]
    vals = [v for v in (r[i] for r in t.rows for i, _, k in shown if k in _NUMERIC) if isinstance(v, (int, float))]
    top = max(vals) if t.heat and vals else 0
    for ci, (i, _, kind) in enumerate(shown):
        if kind in _NUMERIC:
            cmds.append(("ALIGN", (ci, 1), (ci, -1), "RIGHT"))
        if kind == "bool":
            cmds.append(("ALIGN", (ci, 1), (ci, -1), "CENTER"))
        for ri, r in enumerate(t.rows, start=1):
            v = r[i]
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                continue
            if t.heat and kind in _NUMERIC and top > 0:
                a = max(0.0, min(1.0, v / top))
                cmds.append(("BACKGROUND", (ci, ri), (ci, ri),
                             rl_colors.Color(1 - 0.05 * a, 1 - 0.35 * a, 1 - 0.22 * a)))
            elif kind == "pct" and v:
                cmds.append(("TEXTCOLOR", (ci, ri), (ci, ri), _hex(GREEN if v < 0 else RED)))
                cmds.append(("FONTNAME", (ci, ri), (ci, ri), FONT_BOLD))
    tbl = Table(data, repeatRows=1, colWidths=widths, hAlign="LEFT")
    tbl.setStyle(TableStyle(cmds))
    return tbl


def _pdf_section(t: _Table, usable: float) -> list:
    caption = "  ·  ".join(x for x in (f"{len(t.rows):,} rows", t.note) if x)
    head = [Paragraph(escape(t.title), _St.h2), Paragraph(escape(caption), _St.note)]
    png = _table_chart_png(t, usable / cm, 6.2)
    if png:
        head += [RLImage(png, width=usable, height=6.2 * cm), Spacer(1, 6)]
    table = _pdf_table(t, usable)
    if len(t.rows) <= 8:   # short tables never split from their heading
        return [KeepTogether(head + [table]), Spacer(1, 6)]
    # Longer ones start on this page when a few rows fit; the header row repeats.
    return [CondPageBreak((5 if png is None else 11) * cm), KeepTogether(head), table, Spacer(1, 6)]


def _pdf_notes(notes: list[str]) -> list:
    if not notes:
        return []
    return [CondPageBreak(3 * cm), Paragraph("About this data", _St.h2), Spacer(1, 3)] + [
        Paragraph(escape(n), _St.bullet, bulletText="•") for n in notes
    ]


def _report_pdf(rep: _Report) -> io.BytesIO:
    buf = io.BytesIO()
    wide = any(len([c for c in t.columns if c[1] != "url"]) > 7 for t in rep.tables)
    page = landscape(A4) if wide else A4
    doc = _pdf_doc(buf, page, rep.title)
    usable = page[0] - 3.0 * cm
    story: list = [
        Paragraph(escape(rep.title), _St.title),
        Paragraph(escape(rep.subtitle), _St.sub),
    ]
    if rep.kpis:
        for i in range(0, len(rep.kpis), 5):
            story += [_pdf_kpis(rep.kpis[i:i + 5], usable), Spacer(1, 6)]
    for t in rep.tables:
        story += _pdf_section(t, usable)
    story += _pdf_notes(rep.notes)
    doc.build(story, canvasmaker=_canvas_maker(rep.title, rep.stamp))
    return buf


# ── CSV and JSON ─────────────────────────────────────────────────────────────


def _csv_header(label: str, kind: str) -> str:
    unit = _unit(kind)
    return f"{label} ({unit})" if unit else label


def _csv_one(t: _Table) -> str:
    out = io.StringIO()
    w = csv.writer(out, lineterminator="\r\n")
    w.writerow([_csv_header(c, k) for c, k in t.columns])
    for r in t.rows:
        w.writerow(["" if v is None else ("yes" if v is True else "no" if v is False else _plain(v, k))
                    for v, (_, k) in zip(r, t.columns)])
    # BOM so Excel opens UTF-8 (brand names with accents) correctly.
    return "﻿" + out.getvalue()


def _report_csv(rep: _Report, base: str) -> tuple[io.BytesIO, str, str]:
    """One table: a CSV. Several: a ZIP with a CSV per table and a README."""
    if len(rep.tables) == 1:
        return io.BytesIO(_csv_one(rep.tables[0]).encode("utf-8")), f"{base}.csv", _MEDIA["csv"]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        files = []
        for i, t in enumerate(rep.tables, start=1):
            name = f"{i:02d}_{_key(t.title)}.csv"
            files.append(f"  {name:<52} {t.title} ({len(t.rows)} rows)")
            z.writestr(name, _csv_one(t))
        readme = [f"MarketPulse · {rep.title}", rep.subtitle, f"Generated {rep.stamp} · {SOURCE}", "",
                  "Files:", *files, "", "Format: UTF-8 with BOM, comma-separated, dot decimals,",
                  "dates as YYYY-MM-DD, times in UTC. Units are in the column headers.", ""]
        if rep.notes:
            readme += ["About this data:", *[f"  - {n}" for n in rep.notes]]
        z.writestr("README.txt", "\r\n".join(readme))
    return buf, f"{base}.zip", "application/zip"


def _kpi_value(text: str):
    """"141.00 €" -> 141.0, "−1.4%" -> -1.4, "25,878" -> 25878; other text stays text."""
    raw = text.replace(" ", "").replace(",", "").replace("−", "-").replace("€", "").replace("%", "").strip()
    try:
        return int(raw) if raw.lstrip("+-").isdigit() else float(raw)
    except ValueError:
        return text


def _report_json(rep: _Report) -> io.BytesIO:
    doc = {
        "schema_version": SCHEMA_VERSION,
        "report": rep.title,
        "description": rep.subtitle,
        "generated_at": rep.generated.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "source": f"https://{SOURCE}",
        "currency": "EUR",
        "summary": [{"key": _key(k[0]), "label": k[0], "value": _kpi_value(k[1]), "display": k[1]} for k in rep.kpis],
        "notes": rep.notes,
        "tables": [
            {
                "id": _key(t.title),
                "title": t.title,
                "row_count": len(t.rows),
                "columns": [{"key": _key(c), "label": c, "type": k, **({"unit": _unit(k)} if _unit(k) else {})}
                            for c, k in t.columns],
                "rows": [{_key(c): _plain(_as_date(v) if k == "date" else v, k)
                          for v, (c, k) in zip(r, t.columns)} for r in t.rows],
            }
            for t in rep.tables
        ],
    }
    return io.BytesIO(json.dumps(doc, ensure_ascii=False, indent=2, default=str).encode("utf-8"))


_MEDIA = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "csv": "text/csv; charset=utf-8",
    "json": "application/json",
}


def _render(rep: _Report, format: str, base: str) -> StreamingResponse:
    if format == "csv":
        buf, name, media = _report_csv(rep, base)
        return _stream(buf, name, media)
    if format == "json":
        buf = _report_json(rep)
    elif format == "pdf":
        buf = _report_pdf(rep)
    else:
        buf = _report_xlsx(rep)
    return _stream(buf, f"{base}.{format}", _MEDIA[format])


# ── Single product ───────────────────────────────────────────────────────────


@dataclass
class _ProductFacts:
    current: Optional[float]
    previous: Optional[float]
    change: Optional[float]
    low: Optional[float]
    high: Optional[float]
    avg: Optional[float]
    records: int
    cheapest: Optional[Product]


def _product_facts(product: Product, history: list, offers: list[Product]) -> _ProductFacts:
    current = _num(product.last_price) or None
    previous = _num(product.old_price) or None
    prices = [float(h.price) for h in history if h.price]
    priced = [o for o in offers if o.last_price]
    return _ProductFacts(
        current=current,
        previous=previous if previous != current else None,
        change=_delta_percent(current, previous),
        low=min(prices) if prices else current,
        high=max(prices) if prices else current,
        avg=round(sum(prices) / len(prices), 2) if prices else current,
        records=len(prices),
        cheapest=min(priced, key=lambda o: float(o.last_price)) if priced else None,
    )


def _product_report(product: Product, history: list, offers: list[Product], period: str) -> tuple[_Report, _ProductFacts]:
    f = _product_facts(product, history, offers)
    low_offer = float(f.cheapest.last_price) if f.cheapest else None
    shop_prices = _Table("Shop prices", [
        ("Shop", "text"), ("Title", "text"), ("Price", "eur"), ("Previous", "eur"), ("Change", "pct"),
        ("vs cheapest", "share"), ("Promo price", "eur"), ("Promo code", "text"), ("In stock", "bool"),
        ("URL", "url"),
    ], [
        [store_label(store_key(o.search_term)), o.title, _num(o.last_price), _num(o.old_price),
         _delta_percent(_num(o.last_price), _num(o.old_price)),
         round((float(o.last_price) - low_offer) / low_offer * 100, 1) if o.last_price and low_offer else None,
         _num(o.promo_price), o.promo_code, o.is_available, o.url]
        for o in offers
    ], note=f"Same perfume and volume, in stock today; {len(offers)} shop{'s' if len(offers) != 1 else ''}")
    hist = _Table("Price history", [
        ("Date", "datetime"), ("Price", "eur"), ("In stock", "bool"), ("Event", "text"),
    ], [
        [h.parsed_at, _num(h.price), h.in_stock, CHANGE_LABEL.get(h.change_type, h.change_type)]
        for h in history
    ], note=f"Every recorded check at {store_label(store_key(product.search_term))}, oldest first",
        chart=("line", [1]))

    kpis = [
        ("Current price", _eur(f.current)),
        ("Change", _pct(f.change) if f.change is not None else "no change", _tone(f.change)),
        ("Lowest in period", _eur(f.low)),
        ("Highest in period", _eur(f.high)),
        ("Average in period", _eur(f.avg)),
    ]
    rep = _Report(
        # "Bvlgari Man Wood Essence" already starts with the brand.
        title=(product.title if product.brand and (product.title or "").lower().startswith(product.brand.lower())
               else f"{product.brand or ''} {product.title}").strip(),
        subtitle=" · ".join(x for x in [
            store_label(store_key(product.search_term)), product.volume, product.concentration,
            f"Period: {PERIOD_LABEL.get(period, period)}"] if x),
        tables=[shop_prices, hist],
        kpis=kpis,
        notes=[
            "Prices are regular shelf prices in EUR as shown by the shop when it was last checked.",
            "Promo price is the price with the shop's promo code or member price (e.g. 220.lv PLUS); "
            "it is not included in Lowest/Highest.",
            "Change compares the current price with the previous one the shop showed.",
            "vs cheapest is how much dearer a shop is than the cheapest shop for the same bottle.",
        ],
    )
    return rep, f


def _build_product_xlsx(product: Product, rep: _Report, f: _ProductFacts, period: str) -> io.BytesIO:
    wb = _xl_workbook(rep)
    s = wb.active
    s.title = "Summary"
    s.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", (19, 19, 19, 19, 19, 3, 22)):
        s.column_dimensions[col].width = w

    s["A1"] = "PRICE REPORT"
    s["A1"].font = Font(size=9, bold=True, color=ACCENT)
    s["A2"] = (product.brand or "—").upper()
    s["A2"].font = Font(size=11, bold=True, color=GRAY)
    s["A3"] = product.title or "—"
    s["A3"].font = Font(size=20, bold=True, color=INK)
    s.row_dimensions[3].height = 30
    s["A4"] = rep.subtitle
    s["A4"].font = Font(size=10, color=GRAY)

    for i, k in enumerate(rep.kpis):
        col = i + 1
        lab = s.cell(row=6, column=col, value=k[0].upper())
        lab.font = Font(size=8, bold=True, color=GRAY)
        tone = k[2] if len(k) > 2 else None
        num = {0: f.current, 2: f.low, 3: f.high, 4: f.avg}.get(i)
        val = s.cell(row=7, column=col, value=num if num is not None else k[1])
        if num is not None:
            val.number_format = _XL_FMT["eur"]
        val.font = Font(size=17, bold=True, color=GREEN if tone == "down" else RED if tone == "up" else INK)
        for rr in (6, 7):
            c = s.cell(row=rr, column=col)
            c.fill = PatternFill("solid", fgColor=LIGHT)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    s.row_dimensions[6].height = 18
    s.row_dimensions[7].height = 32

    s["A9"] = "Details"
    s["A9"].font = Font(size=12, bold=True, color=INK)
    shops = len(rep.tables[0].rows)
    cheapest = f.cheapest
    details = [
        ("Shop", store_label(store_key(product.search_term))),
        ("Volume · type", " · ".join(x for x in (product.volume, product.concentration) if x) or "—"),
        ("In stock", "Yes" if product.is_available else "No"),
        ("Previous price", f.previous or "—"),
        ("Promo price", _num(product.promo_price) or "—"),
        ("Promo code", product.promo_code or "—"),
        ("Comparison", f"Cheapest of {shops} shop{'s' if shops != 1 else ''}: "
                       f"{store_label(store_key(cheapest.search_term))}, {_eur(_num(cheapest.last_price))}"
         if cheapest else "Only this shop"),
        ("Price records", f"{f.records} ({PERIOD_LABEL.get(period, period).lower()})"),
        ("Last checked", product.last_seen.strftime("%Y-%m-%d %H:%M UTC") if product.last_seen else "—"),
        ("Product page", product.url or "—"),
        ("Generated", rep.stamp),
    ]
    for i, (k, v) in enumerate(details, start=10):
        a = s.cell(row=i, column=1, value=k)
        a.font = Font(size=10, color=GRAY)
        s.merge_cells(start_row=i, start_column=2, end_row=i, end_column=5)
        b = s.cell(row=i, column=2, value=v)
        b.font = Font(size=10, color=INK)
        if isinstance(v, float):
            b.number_format = _XL_FMT["eur"]
            b.alignment = Alignment(horizontal="left")
        for ci in range(1, 6):
            s.cell(row=i, column=ci).border = Border(bottom=_HAIR)
        if k == "Product page" and product.url:
            b.hyperlink = product.url
            b.value = "Open on the shop's site ↗"
            b.font = Font(size=10, color=LINK, underline="single")
    r = 10 + len(details) + 1
    s.cell(row=r, column=1, value="Sheets: Shop prices · Price history (with chart)").font = Font(size=9, italic=True, color=GRAY)

    img_bytes = _fetch_image_bytes(product.image_url)
    try:
        img = XLImage(io.BytesIO(img_bytes) if img_bytes else _placeholder_png(product.brand))
        ratio = (img.height / img.width) if img.width else 1
        img.width, img.height = 150, int(150 * ratio) if ratio < 1.6 else 240
        if ratio >= 1.6:
            img.width = int(240 / ratio)
        s.add_image(img, "G1")
    except Exception:
        pass
    _xl_print(s, None, rep.title, landscape_page=False)

    for i, t in enumerate(rep.tables, start=1):
        ws = wb.create_sheet(t.title)
        _xl_head(ws, t.title, f"{rep.title}  ·  {t.note}")
        last = _xl_table(ws, t, 4, f"tbl{i}_{_key(t.title)[:40]}")
        if t.title == "Price history" and len(t.rows) >= 2:
            ch = LineChart()
            ch.height, ch.width = 9, 20
            ch.legend = None
            ch.y_axis.delete = False
            ch.x_axis.delete = False
            ch.y_axis.number_format = '#,##0 "€"'
            ch.x_axis.number_format = "dd.mm"
            _xl_fit_axis(ch, [r[1] for r in t.rows])
            ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=last), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last))
            ch.series[0].graphicalProperties.line.solidFill = INK
            ch.series[0].graphicalProperties.line.width = 22000
            ch.series[0].smooth = False
            ws.add_chart(ch, "F4")
        _xl_print(ws, 4, rep.title)

    buf = io.BytesIO()
    wb.save(buf)
    return buf


def _build_product_pdf(product: Product, history: list, rep: _Report, f: _ProductFacts) -> io.BytesIO:
    buf = io.BytesIO()
    page = A4
    doc = _pdf_doc(buf, page, rep.title)
    usable = page[0] - 3.0 * cm

    img_bytes = _fetch_image_bytes(product.image_url)
    img = RLImage(io.BytesIO(img_bytes) if img_bytes else _placeholder_png(product.brand),
                  width=3.6 * cm, height=3.6 * cm, kind="proportional")
    price_line = f'<font size="22"><b>{escape(_eur(f.current))}</b></font>'
    if f.previous:
        price_line += f'   <font color="#{MUTED}"><strike>{escape(_eur(f.previous))}</strike></font>'
    if f.change is not None:
        price_line += f'   <font color="#{GREEN if f.change < 0 else RED}"><b>{escape(_pct(f.change))}</b></font>'
    promo = ""
    if product.promo_price:
        promo = (f'<font color="#{ACCENT}"><b>{escape(_eur(_num(product.promo_price)))}</b> with '
                 f'{escape(product.promo_code or "promo")}</font>')
    info = [
        Paragraph(f'<font color="#{ACCENT}"><b>{escape((product.brand or "—").upper())}</b></font>', _St.body),
        Paragraph(escape(product.title or "—"), ParagraphStyle("pt", parent=_St.title, fontSize=17, leading=21)),
        Paragraph(escape(rep.subtitle), _St.sub),
        Paragraph(price_line, ParagraphStyle("pl", parent=_St.body, leading=26)),
    ]
    if promo:
        info.append(Paragraph(promo, _St.body))
    if product.url:
        info.append(Spacer(1, 3))
        info.append(Paragraph(f'<link href="{escape(product.url)}"><font color="#{LINK}">'
                              f'Open on the shop\'s site ↗</font></link>', _St.note))
    hero = Table([[img, info]], colWidths=[4.1 * cm, usable - 4.1 * cm])
    hero.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0),
                              ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
    story: list = [hero, Spacer(1, 12), _pdf_kpis(rep.kpis, usable), Spacer(1, 4)]

    story += [Paragraph("Price history", _St.h2),
              Paragraph(escape(f"{f.records} price records · {rep.subtitle.split(' · ')[-1]}"), _St.note),
              RLImage(_history_png([(h.parsed_at, float(h.price)) for h in history], usable / cm, 6.4),
                      width=usable, height=6.4 * cm)]

    shop_prices, hist = rep.tables
    story += _pdf_section(shop_prices, usable)
    events = [r for r in hist.rows if r[3] != CHANGE_LABEL["no_change"]][-15:][::-1]
    story += _pdf_section(_Table("Recent price changes", hist.columns, events,
                                 note="Newest first; checks with no change are left out"), usable)
    story += _pdf_notes(rep.notes)
    doc.build(story, canvasmaker=_canvas_maker(rep.title, rep.stamp))
    return buf


# ── Endpoints ─────────────────────────────────────────────────────────────────


@router.get("/products/{product_id}")
async def export_product(
    product_id: int,
    format: str = Query("xlsx", pattern=_FORMATS),
    period: str = Query("3m", pattern="^(1w|1m|3m|6m|1y|all)$"),
    db: AsyncSession = Depends(get_db),
):
    """Export a single product: price history and today's shop prices."""
    stmt = (
        select(Product)
        .options(selectinload(Product.history))
        .where(Product.id == product_id)
    )
    product = (await db.execute(stmt)).scalar_one_or_none()
    if not product:
        raise HTTPException(404, "Product not found")

    cutoff = _period_cutoff(period)
    history = [h for h in product.history if not cutoff or h.parsed_at >= cutoff]
    history.sort(key=lambda h: h.parsed_at)

    offers: list[Product] = []
    if product.group_id:
        norm_vol = (product.volume or "").lower().replace(" ", "").strip()
        grp_stmt = select(Product).where(
            Product.group_id == product.group_id,
            Product.is_available == True,  # noqa: E712
            Product.last_price.isnot(None),
            Product.last_price > 0,
        )
        grp_products = (await db.execute(grp_stmt)).scalars().all()
        by_store: dict[str, Product] = {}
        for p in grp_products:
            if norm_vol and (p.volume or "").lower().replace(" ", "").strip() != norm_vol:
                continue
            s_key = (p.search_term or "").lower().strip()
            if s_key not in by_store or float(p.last_price) < float(by_store[s_key].last_price):
                by_store[s_key] = p
        offers = sorted(by_store.values(), key=lambda p: float(p.last_price))[:5]
    if not offers:
        offers = [product]

    rep, facts = _product_report(product, history, offers, period)
    base = _safe_filename(rep.title.replace(" ", "_"))
    if format == "xlsx":
        return _stream(_build_product_xlsx(product, rep, facts, period), f"{base}.xlsx", _MEDIA["xlsx"])
    if format == "pdf":
        return _stream(_build_product_pdf(product, history, rep, facts), f"{base}.pdf", _MEDIA["pdf"])
    return _render(rep, format, base)


# ── Catalogue ─────────────────────────────────────────────────────────────────


def _filters_desc(**f) -> str:
    flags = {"in_stock": "in stock", "promo": "with a promo price", "multi_shop": "in 2+ shops",
             "dropped": "price dropped"}
    names = {"price_min": "price from", "price_max": "price up to", "types": "type", "stores": "shops"}
    parts = []
    for k, v in f.items():
        if v in (None, [], False, ""):
            continue
        if v is True and k in flags:
            parts.append(flags[k])
            continue
        if k in ("price_min", "price_max"):
            v = _eur(v)
        elif k == "stores":
            v = [store_label(store_key(s)) for s in v]
        parts.append(f"{names.get(k, k.replace('_', ' '))}: {', '.join(map(str, v)) if isinstance(v, list) else v}")
    return "  ·  ".join(parts) or "all perfumes"


@router.get("/catalog")
async def export_catalog(
    format: str = Query("xlsx", pattern=_FORMATS),
    view: str = Query("cards", pattern="^(cards|listings)$"),
    keys: List[str] = Query(None, description="Selected perfume cards (view=cards)"),
    ids: List[int] = Query(None, description="Selected shop listings (view=listings)"),
    search: Optional[str] = None,
    brands: List[str] = Query(None),
    stores: List[str] = Query(None),
    volumes: List[str] = Query(None),
    types: List[str] = Query(None),
    price_min: Optional[float] = None,
    price_max: Optional[float] = None,
    in_stock: Optional[bool] = None,
    promo: Optional[bool] = None,
    multi_shop: Optional[bool] = None,
    dropped: Optional[bool] = None,
    sort: Optional[str] = None,
    limit: int = Query(2000, ge=1, le=5000),
    db: AsyncSession = Depends(get_db),
):
    """Catalogue export: the selected rows, or everything matching the filters.

    view=cards writes one row per perfume and size with a price column per
    shop; view=listings writes one row per shop listing, like the table view.
    """
    filters = dict(search=search, brands=brands, stores=stores, volumes=volumes, types=types,
                   price_min=price_min, price_max=price_max, in_stock=in_stock, promo=promo)
    stamp = datetime.utcnow().strftime("%Y%m%d")
    desc = f"{len(keys or ids or [])} selected" if (keys or ids) else _filters_desc(
        **filters, multi_shop=multi_shop, dropped=dropped)

    if view == "cards":
        if keys:
            cards = await cards_for_keys(db, keys[:limit])
        else:
            page = await query_cards(db, **filters, multi_shop=multi_shop, dropped=dropped,
                                     sort=sort if sort in ("popular", "shops", "drop", "low", "high", "gap", "brand", "new") else "popular",
                                     skip=0, limit=limit)
            cards = page.items
        shops = [k for k in STORE_LABELS if any(o.store == k for c in cards for o in c.offers)]
        cols = [("Brand", "text"), ("Perfume", "text"), ("Type", "text"), ("Volume", "text"),
                ("Cheapest shop", "text"), ("Lowest", "eur"), ("Highest", "eur"), ("Gap", "share"), ("Shops", "int"),
                ("30-day change", "pct")] + [(f"{store_label(k)}", "eur") for k in shops]
        rows = []
        for c in cards:
            by_shop = {}
            for o in c.offers:
                if o.in_stock and (o.store not in by_shop or o.price < by_shop[o.store]):
                    by_shop[o.store] = o.price
            gap = round((c.max_price - c.min_price) / c.min_price * 100, 1) if c.min_price and c.max_price else None
            rows.append([c.brand, c.name, c.concentration, c.volume, store_label(c.cheapest_store),
                         c.min_price, c.max_price, gap, c.shops, c.change_30d] + [by_shop.get(k) for k in shops])
        table = _Table("Perfumes", cols, rows, note="One row per perfume and size; shop columns show in-stock prices")
        multi = [r for r in rows if (r[8] or 0) >= 2]
        gaps = sorted(r[7] for r in multi if r[7] is not None)
        wins = {k: sum(1 for c in cards if c.shops >= 2 and c.cheapest_store == k) for k in shops}
        best = max(wins, key=wins.get, default=None)
        kpis = [
            ("Perfumes", f"{len(rows):,}"),
            ("In 2+ shops", f"{len(multi):,}"),
            ("Median gap", f"{gaps[len(gaps) // 2]:.1f}%" if gaps else "—"),
            ("Cheapest most often", f"{store_label(best)} ({wins[best]})" if best and wins[best] else "—"),
        ]
        notes = [
            "Lowest and Highest are the cheapest and dearest in-stock shop prices for the same perfume and size.",
            "Gap is how much dearer the dearest shop is than the cheapest one.",
            "30-day change compares the lowest price now with the lowest price 30 days ago.",
            "An empty shop column means the shop does not sell that size, or it is out of stock.",
        ]
    else:
        page = await query_listings(db, ids=ids, **filters, multi_shop=multi_shop, dropped=dropped,
                                    sort=sort if sort in ("popular", "brand", "name", "price", "change", "shop") else "popular",
                                    skip=0, limit=limit)
        rows = [
            [r.brand, r.name, r.concentration, r.volume, r.store_name, r.price, r.old_price, r.promo_price,
             r.promo_code, r.change_30d, r.in_stock, r.checked_at, r.url]
            for r in page.items
        ]
        table = _Table("Shop listings", [
            ("Brand", "text"), ("Perfume", "text"), ("Type", "text"), ("Volume", "text"), ("Shop", "text"),
            ("Price", "eur"), ("Previous", "eur"), ("Promo price", "eur"), ("Promo code", "text"),
            ("30-day change", "pct"), ("In stock", "bool"), ("Checked", "datetime"), ("URL", "url"),
        ], rows, note="One row per shop listing")
        kpis = [
            ("Listings", f"{len(rows):,}"),
            ("Shops", f"{len({r[4] for r in rows}):,}"),
            ("In stock", f"{sum(1 for r in rows if r[10]):,}"),
            ("With promo price", f"{sum(1 for r in rows if r[7]):,}"),
        ]
        notes = [
            "Price is the regular shelf price; Promo price needs the shop's promo code or membership.",
            "Previous is the price the shop showed before its latest change.",
            "Checked is the last time MarketPulse read the listing (UTC).",
        ]
    rep = _Report("Catalogue", f"Filters: {desc}", [table], kpis=kpis, notes=notes)
    return _render(rep, format, f"marketpulse_catalogue_{stamp}")


# ── Analytics ────────────────────────────────────────────────────────────────

_SECTIONS = ("overview", "shops", "drops", "rises", "gaps", "brands", "bands", "activity", "index")


@router.get("/analytics")
async def export_analytics(
    format: str = Query("xlsx", pattern=_FORMATS),
    days: int = Query(7, ge=1, le=90, description="Window for the biggest moves"),
    section: Optional[str] = Query(None, description="Only this section (CSV of one table)"),
    db: AsyncSession = Depends(get_db),
):
    """Everything on the analytics page, one table per section."""
    if section and section not in _SECTIONS:
        raise HTTPException(422, f"section must be one of {', '.join(_SECTIONS)}")
    want = (lambda name: section is None or section == name)
    snap = await market.snapshot(db)
    tables: list[_Table] = []
    if want("overview"):
        tables.append(_Table("Market overview", [("Metric", "text"), ("Value", "text")], [
            ["Latest scan", snap.as_of.strftime("%Y-%m-%d %H:%M UTC") if snap.as_of else "—"],
            ["Listings tracked", f"{snap.total_products:,}"],
            ["In stock now", f"{snap.current_products:,}"],
            ["Bottles in 2+ shops", f"{snap.comparable:,}"],
            ["Price changes, 7 days", f"{snap.changes_7d:,}"],
            ["Price changes, 30 days", f"{snap.changes_30d:,}"],
            ["Median gap, dearest vs cheapest shop", f"{snap.spread_median:.1f}%"],
            ["Mean gap", f"{snap.spread_mean:.1f}%"],
            ["Price records since", snap.history_since.strftime("%Y-%m-%d") if snap.history_since else "—"],
        ]))
    if want("shops"):
        tables.append(_Table("Shops", [
            ("Shop", "text"), ("Listings", "int"), ("In stock", "int"), ("Changes 7d", "int"),
            ("Changes 30d", "int"), ("Cheapest for", "int"), ("Cheapest share", "share"), ("Promo listings", "int"),
        ], [[s.store_name, s.products, s.current, s.changes_7d, s.changes_30d, s.cheapest_count,
             s.cheapest_share, s.promo_listings] for s in snap.shops],
            note="Cheapest for: bottles where the shop has the lowest price", chart=("bar", [2, 5])))
    move_cols = [("Brand", "text"), ("Perfume", "text"), ("Type", "text"), ("Volume", "text"), ("Shop", "text"),
                 ("Before", "eur"), ("Now", "eur"), ("Change", "pct")]
    if want("drops"):
        rows = await market.moves(db, days=days, direction="down", limit=30)
        tables.append(_Table(f"Biggest drops, {days} days", move_cols, [
            [m.brand, m.name, m.concentration, m.volume, m.store_name, m.old_price, m.new_price, m.change_percent]
            for m in rows], note="Top 30 by percent"))
    if want("rises"):
        rows = await market.moves(db, days=days, direction="up", limit=30)
        tables.append(_Table(f"Biggest rises, {days} days", move_cols, [
            [m.brand, m.name, m.concentration, m.volume, m.store_name, m.old_price, m.new_price, m.change_percent]
            for m in rows], note="Top 30 by percent"))
    if want("gaps"):
        cards = await cards_for_keys(db, await market.top_gap_keys(db, limit=50))
        rows = []
        for c in cards:
            live = [o for o in c.offers if o.in_stock] or c.offers
            dear = max(live, key=lambda o: o.price)
            rows.append([c.brand, c.name, c.concentration, c.volume, store_label(c.cheapest_store), c.min_price,
                         dear.store_name, dear.price, round((dear.price - c.min_price) / c.min_price * 100, 1)])
        rows.sort(key=lambda r: -r[-1])
        tables.append(_Table("Widest gaps between shops", [
            ("Brand", "text"), ("Perfume", "text"), ("Type", "text"), ("Volume", "text"),
            ("Cheapest shop", "text"), ("Lowest", "eur"), ("Dearest shop", "text"), ("Highest", "eur"), ("Gap", "share"),
        ], rows, note="Same bottle, in stock; top 50 by gap"))
    if want("brands"):
        m = await market.brand_shop(db, top=20)
        cell = {(c.brand, c.store): c for c in m.cells}
        tables.append(_Table("Brand × shop premium over the cheapest", [("Brand", "text")] + [
            (store_label(s), "share") for s in m.stores
        ], [[b] + [cell[(b, s)].premium if (b, s) in cell else None for s in m.stores] for b in m.brands],
            note="Average % above the cheapest shop; 0% = usually the cheapest; 20 largest brands", heat=True))
    if want("bands"):
        bands = await market.price_bands(db)
        labels = list(dict.fromkeys(b.label for b in sorted(bands, key=lambda b: b.band)))
        stores = list(dict.fromkeys(b.store for b in bands))
        cnt = {(b.store, b.label): b.count for b in bands}
        tables.append(_Table("In-stock listings by price band", [("Shop", "text")] + [(lb, "int") for lb in labels],
                             [[store_label(s)] + [cnt.get((s, lb), 0) for lb in labels] for s in stores],
                             note="Number of listings per price range", heat=True))
    if want("activity"):
        act = await market.activity(db, 30)
        tables.append(_Table("Price changes per day, 30 days", [
            ("Date", "date"), ("Shop", "text"), ("Rises", "int"), ("Drops", "int"),
        ], [[a.date, store_label(a.store), a.up, a.down] for a in act]))
    if want("index"):
        idx = await market.price_index(db, 30)
        shops = list(idx.shops)
        tables.append(_Table("Price index, 100 = 30 days ago", [("Date", "date")] + [
            (store_label(s), "num") for s in shops
        ], [[d] + [idx.shops[s][i] for s in shops] for i, d in enumerate(idx.days)],
            note="Average price of the same bottles; below 100 = cheaper than 30 days ago",
            chart=("line", list(range(1, len(shops) + 1)))))

    kpis = [
        ("Listings tracked", f"{snap.total_products:,}"),
        ("In stock now", f"{snap.current_products:,}"),
        ("Bottles in 2+ shops", f"{snap.comparable:,}"),
        (f"Price changes, 7 days", f"{snap.changes_7d:,}"),
        ("Median gap", f"{snap.spread_median:.1f}%"),
    ]
    rep = _Report(
        "Market analytics",
        f"Moves over {days} days  ·  latest scan "
        f"{snap.as_of.strftime('%Y-%m-%d %H:%M UTC') if snap.as_of else '—'}",
        tables, kpis=kpis,
        notes=[
            "Covers men's perfumes in Notino, Makeup.lv, 220.lv and Douglas; prices in EUR as listed.",
            "Gap = how much dearer the dearest shop is than the cheapest for the same perfume and size.",
            "Price index: 100 is each shop's average price for the same bottles 30 days ago.",
            "Price records start on "
            f"{snap.history_since.strftime('%Y-%m-%d') if snap.history_since else '—'}; "
            "windows are at most 90 days.",
        ],
    )
    stamp = datetime.utcnow().strftime("%Y%m%d")
    return _render(rep, format, f"marketpulse_analytics_{stamp}")
